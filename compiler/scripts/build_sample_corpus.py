"""Build the sample corpus in data/raw/: BürgerEnergie Eschenbrück eG.

A fictional citizen energy cooperative in the fictional village of
Eschenbrück builds a rooftop solar plant with battery storage on the local
primary school ("Sonnendach Lindenhof"). The sources are what such a
project really produces: a grant application, board minutes, a structural
survey, installer and grid-operator emails, a budget workbook, member
slides, a public flyer and FAQ, a newspaper article, monitoring exports
and an invoice archive -- in English and German, in 15 file formats.

The facts were chosen to exercise the compiler: several values change over
time (plant size, module type, battery size, commissioning date, member
count, heat-pump pilot date), and some sources are stale or plainly wrong
(the March flyer's share price, the FAQ's battery size). The ground truth
is documented in documentation/18-sample-domain.md and labelled in
data/trust_eval_dataset.json and data/qa_benchmark.json.

Every person, company, place and number here is invented.

    python scripts/build_sample_corpus.py            # (re)write data/raw/
    python scripts/build_sample_corpus.py --out DIR  # somewhere else
    python scripts/build_sample_corpus.py --clean    # empty data/raw/ first

The output is deterministic: office files and archives get fixed
timestamps, so re-running the script does not change committed files.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import shutil
import sys
import textwrap
import zipfile
from datetime import datetime
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

COMPILER_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(COMPILER_DIR))

from models import RAW_DIR  # noqa: E402

FIXED_TIME = datetime(2026, 9, 15, 12, 0, 0)
ZIP_TIME = (2026, 9, 15, 12, 0, 0)
DOMAIN = "eschenbrueck-energie.example"


# --- Writers -------------------------------------------------------------------


def _normalize_zip(path: Path) -> None:
    """Rewrite a zip-based file (docx/xlsx/pptx/zip) with fixed entry
    timestamps and order, so its bytes don't depend on when it was built."""
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info.filename)) for info in src.infolist()]
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    # openpyxl writes the save time into core.xml whatever `modified` says.
    entries = [
        (name, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, data) if name == "docProps/core.xml" else data)
        for name, data in entries
    ]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def _pdf_string(text: str) -> str:
    """A PDF literal string in WinAnsiEncoding (cp1252), ASCII-safe."""
    out = []
    for byte in text.encode("cp1252", errors="replace"):
        char = chr(byte)
        if char in "()\\":
            out.append("\\" + char)
        elif 32 <= byte < 127:
            out.append(char)
        else:
            out.append(f"\\{byte:03o}")
    return "(" + "".join(out) + ")"


def write_pdf(path: Path, title: str, blocks: list[tuple[str, str]]) -> None:
    """A minimal text PDF (Helvetica, A4, wrapped lines) that pypdf can
    extract. blocks: (style, text), style in {"h1", "h2", "p", "row"}."""
    lines: list[tuple[str, int, str]] = []  # (font, size, text)
    for style, text in blocks:
        if style == "h1":
            lines += [("F2", 16, text), ("F1", 11, "")]
        elif style == "h2":
            lines += [("F1", 11, ""), ("F2", 12, text)]
        elif style == "row":
            lines.append(("F1", 10, text))
        else:
            for wrapped in textwrap.wrap(text, 92) or [""]:
                lines.append(("F1", 10, wrapped))
            lines.append(("F1", 10, ""))
    per_page = 52
    pages = [lines[i : i + per_page] for i in range(0, len(lines), per_page)] or [[]]

    objects: list[bytes] = []

    def add(obj: str | bytes) -> int:
        objects.append(obj.encode("latin-1") if isinstance(obj, str) else obj)
        return len(objects)

    catalog = add("")  # placeholder, filled in below
    pages_obj = add("")
    font1 = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    font2 = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    page_ids = []
    for page_number, page_lines in enumerate(pages, start=1):
        ops = ["BT", "50 790 Td"]
        for font, size, text in page_lines:
            ops.append(f"/{font} {size} Tf")
            ops.append(f"{_pdf_string(text)} Tj")
            ops.append(f"0 -{size + 4} Td")
        ops.append("ET")
        ops.append(f"BT /F1 8 Tf 50 30 Td {_pdf_string(f'{title} - {page_number}/{len(pages)}')} Tj ET")
        stream = "\n".join(ops).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent {pages_obj} 0 R /MediaBox [0 0 595 842] "
                f"/Resources << /Font << /F1 {font1} 0 R /F2 {font2} 0 R >> >> /Contents {content} 0 R >>"
            )
        )
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_obj} 0 R >>".encode("latin-1")
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objects[pages_obj - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("latin-1")
    info = add(f"<< /Title {_pdf_string(title)} /Producer (build_sample_corpus.py) /CreationDate (D:20260915120000Z) >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root %d 0 R /Info %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        catalog,
        info,
        xref,
    )
    path.write_bytes(bytes(out))


def write_docx(path: Path, title: str, blocks: list[tuple[str, object]]) -> None:
    """blocks: ("h1"|"h2"|"p"|"bullet", str) or ("table", [[cells], ...])."""
    from docx import Document

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = "BürgerEnergie Eschenbrück eG"
    doc.core_properties.created = FIXED_TIME
    doc.core_properties.modified = FIXED_TIME
    doc.core_properties.last_modified_by = "BürgerEnergie Eschenbrück eG"
    doc.core_properties.revision = 1
    for style, content in blocks:
        if style == "h1":
            doc.add_heading(str(content), level=1)
        elif style == "h2":
            doc.add_heading(str(content), level=2)
        elif style == "bullet":
            doc.add_paragraph(str(content), style="List Bullet")
        elif style == "table":
            rows = list(content)  # type: ignore[arg-type]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
        else:
            doc.add_paragraph(str(content))
    doc.save(path)
    _normalize_zip(path)


def write_xlsx(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for column in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(60, width + 2)
    wb.properties.creator = "Tobias Brandt"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.save(path)
    _normalize_zip(path)


def write_pptx(path: Path, slides: list[tuple[str, list[str]]]) -> None:
    from pptx import Presentation

    prs = Presentation()
    prs.core_properties.title = slides[0][0]
    prs.core_properties.author = "BürgerEnergie Eschenbrück eG"
    prs.core_properties.created = FIXED_TIME
    prs.core_properties.modified = FIXED_TIME
    prs.core_properties.revision = 1
    for index, (title, bullets) in enumerate(slides):
        layout = prs.slide_layouts[0 if index == 0 else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0] if bullets else ""
        for bullet in bullets[1:]:
            body.add_paragraph().text = bullet
    prs.save(path)
    _normalize_zip(path)


def write_eml(
    path: Path,
    *,
    date: str,
    sender: str,
    to: str,
    subject: str,
    body: str,
    cc: str | None = None,
    in_reply_to: str | None = None,
) -> None:
    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = to
    if cc:
        msg["Cc"] = cc
    msg["Subject"] = subject
    msg["Date"] = format_datetime(datetime.fromisoformat(date))
    msg["Message-ID"] = f"<{path.stem}@{DOMAIN}>"
    if in_reply_to:
        msg["In-Reply-To"] = f"<{in_reply_to}@{DOMAIN}>"
        msg["References"] = f"<{in_reply_to}@{DOMAIN}>"
    msg.set_content(textwrap.dedent(body).strip() + "\n")
    path.write_bytes(msg.as_bytes())


def write_text(path: Path, text: str) -> None:
    path.write_text(textwrap.dedent(text).lstrip("\n"), encoding="utf-8")


def write_csv(path: Path, rows: list[list[object]], delimiter: str = ",") -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        csv.writer(fh, delimiter=delimiter, lineterminator="\n").writerows(rows)


def write_png(path: Path) -> None:
    """Roof layout sketch: two roofs, module counts, the battery room."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (900, 560), "white")
    d = ImageDraw.Draw(img)
    d.text((30, 20), "Sonnendach Lindenhof - roof layout (planning sketch, May 2026)", fill="black")
    d.rectangle((40, 70, 520, 330), outline="black", width=3)
    d.text((60, 80), "Roof A: Grundschule am Lindenhof", fill="black")
    for row in range(12):
        for col in range(23):
            x, y = 60 + col * 19, 110 + row * 17
            d.rectangle((x, y, x + 15, y + 13), outline="#1f4e79", fill="#9dc3e6")
    d.text((60, 318), "276 modules (Nordlicht NL-430), south-facing, 15 deg", fill="black")
    d.rectangle((560, 70, 860, 330), outline="black", width=3)
    d.text((575, 80), "Roof B: Sporthalle Nord", fill="black")
    for row in range(9):
        for col in range(14):
            if row * 14 + col >= 123:
                break
            x, y = 575 + col * 19, 110 + row * 20
            d.rectangle((x, y, x + 15, y + 15), outline="#1f4e79", fill="#9dc3e6")
    d.rectangle((575, 290, 845, 318), outline="#c00000", width=2)
    d.text((580, 297), "east section: no modules (load reserve)", fill="#c00000")
    d.text((575, 318), "123 modules", fill="black")
    d.rectangle((40, 380, 520, 520), outline="black", width=2)
    d.text((60, 390), "Basement of the school: battery room", fill="black")
    d.text((60, 415), "100 kWh storage, F90 fire door, 3 inverters WR-1..WR-3 on the wall", fill="black")
    d.text((560, 390), "Total: 399 modules x 430 W = 171.6 kWp", fill="black")
    d.text((560, 415), "Grid connection: new transformer station (Netze Mittelland)", fill="black")
    img.save(path, format="PNG", optimize=False)


# --- Content -------------------------------------------------------------------


def build(out: Path) -> list[Path]:
    written: list[Path] = []

    def target(rel: str) -> Path:
        path = out / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        written.append(path)
        return path

    # 1. Board minutes, January (DOCX, EN) -------------------------------------
    write_docx(
        target("meetings/2026-01-20-board-minutes.docx"),
        "Board meeting minutes 20 January 2026",
        [
            ("h1", "BürgerEnergie Eschenbrück eG - Board meeting minutes"),
            ("p", "Date: 20 January 2026, 19:00-21:15, Gemeindehaus Eschenbrück."),
            ("p", "Present: Dr. Hanna Vogt (chair), Tobias Brandt (finance), Selin Aydın (technology), Klaus Reimann (chair of the supervisory board, guest)."),
            ("h2", "1. Membership"),
            ("p", "Tobias reported 388 members holding 5,912 shares at the end of 2025. The share price stays at 250 euros per share; the statutes allow at most 40 shares per member."),
            ("h2", "2. Project idea: Sonnendach Lindenhof"),
            ("p", "Selin presented the idea of a solar plant on the roofs of the Grundschule am Lindenhof and the Sporthalle Nord. First estimate: about 198 kWp with 450 modules of the type Helion H-440, plus a 150 kWh battery in the school basement so the school can use more of its own solar power."),
            ("p", "The municipality (mayor Petra Lindqvist) has signalled that it would lease both roofs to the cooperative for a symbolic rent."),
            ("h2", "3. Funding"),
            ("p", "The board decided unanimously to apply for the state programme KlimaKommunal before the deadline of 15 February 2026. Tobias will talk to the Raiffeisenkasse Talgrund about a loan."),
            ("h2", "4. Heat pump pilot"),
            ("p", "Hanna raised the idea of a heat pump for the outdoor pool (Freibad Eschenbrück) as a second project. Selin will ask members about it in the February survey."),
            ("h2", "Action items"),
            ("bullet", "Selin: draft grant application (due 10 February)."),
            ("bullet", "Tobias: loan talks with Raiffeisenkasse Talgrund."),
            ("bullet", "Hanna: meet mayor Lindqvist about a roof lease."),
        ],
    )

    # 2. Member survey (JSON) ----------------------------------------------------
    survey = {
        "survey": "Mitgliederbefragung 2026 / member survey 2026",
        "organisation": "BürgerEnergie Eschenbrück eG",
        "fielded": {"from": "2026-02-01", "to": "2026-02-21"},
        "members_invited": 388,
        "responses": 214,
        "questions": [
            {
                "id": "q1",
                "text": "Should the cooperative build the Sonnendach Lindenhof solar plant?",
                "results": {"yes": 0.89, "no": 0.04, "undecided": 0.07},
            },
            {
                "id": "q2",
                "text": "Should the cooperative pilot a heat pump for the Freibad Eschenbrück?",
                "results": {"yes": 0.71, "no": 0.12, "undecided": 0.17},
            },
            {
                "id": "q3",
                "text": "Would you buy additional shares for the project?",
                "results": {"yes": 0.46, "maybe": 0.31, "no": 0.23},
            },
            {
                "id": "q4",
                "text": "Dividend or reinvestment?",
                "results": {"prefer dividend": 0.64, "prefer reinvestment": 0.36},
            },
        ],
        "free_text_themes": [
            "Please keep the share price at 250 euros",
            "Worried about noise from the pool heat pump",
            "Offer shares to pupils' parents",
        ],
    }
    target("survey/2026-02-member-survey-results.json").write_text(
        json.dumps(survey, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    # 3. Grant application (PDF, DE) ---------------------------------------------
    write_pdf(
        target("project/2026-02-10-foerderantrag-klimakommunal.pdf"),
        "Förderantrag KlimaKommunal - Sonnendach Lindenhof",
        [
            ("h1", "Förderantrag im Programm KlimaKommunal"),
            ("p", "Antragstellerin: BürgerEnergie Eschenbrück eG, vertreten durch Dr. Hanna Vogt (Vorstandsvorsitzende). Datum: 10. Februar 2026."),
            ("h2", "1. Vorhaben"),
            ("p", "Die Genossenschaft plant eine Photovoltaikanlage mit 198 kWp auf den Dächern der Grundschule am Lindenhof und der Sporthalle Nord. Vorgesehen sind 450 Module vom Typ Helion H-440 sowie ein Batteriespeicher mit 150 kWh im Keller der Schule."),
            ("p", "Der Strom wird vorrangig an die Schule geliefert; Überschüsse werden in das Netz der Netze Mittelland GmbH eingespeist. Geplante Inbetriebnahme: Juni 2026."),
            ("h2", "2. Klimawirkung"),
            ("p", "Erwarteter Jahresertrag: 190.000 kWh. Erwartete CO2-Einsparung: 72 Tonnen pro Jahr."),
            ("h2", "3. Kosten und Finanzierung"),
            ("row", "Gesamtkosten (brutto):                 312.000 EUR"),
            ("row", "Beantragter Zuschuss KlimaKommunal:     96.000 EUR"),
            ("row", "Darlehen Raiffeisenkasse Talgrund:     130.000 EUR"),
            ("row", "Eigenkapital der Mitglieder:            86.000 EUR"),
            ("p", ""),
            ("h2", "4. Beteiligung der Bürgerinnen und Bürger"),
            ("p", "Die Genossenschaft hat 388 Mitglieder. Ein Geschäftsanteil kostet 250 Euro. Neue Mitglieder aus Eschenbrück können sich ab einem Anteil beteiligen."),
            ("h2", "5. Zeitplan"),
            ("row", "März 2026:   Dachnutzungsvertrag mit der Gemeinde, Netzanschlussanfrage"),
            ("row", "April 2026:  Statikgutachten, Auftragsvergabe"),
            ("row", "Mai 2026:    Montage während der Pfingstferien"),
            ("row", "Juni 2026:   Inbetriebnahme"),
        ],
    )

    # 4. Project plan (DOCX, EN) -------------------------------------------------
    write_docx(
        target("project/2026-03-02-project-plan.docx"),
        "Project plan Sonnendach Lindenhof",
        [
            ("h1", "Project plan: Sonnendach Lindenhof"),
            ("p", "Version 1.0, 2 March 2026. Owner: Selin Aydın (technical lead)."),
            ("h2", "Scope"),
            ("bullet", "Rooftop PV on the Grundschule am Lindenhof (roof A) and the Sporthalle Nord (roof B): 198 kWp, 450 x Helion H-440."),
            ("bullet", "Battery storage: 150 kWh in the school basement."),
            ("bullet", "Power purchase agreement: the school buys the solar power at 19.5 ct/kWh for 20 years."),
            ("bullet", "Out of scope for 2026: heat pump pilot for the Freibad Eschenbrück, planned as a follow-up project for autumn 2026."),
            ("h2", "Milestones"),
            (
                "table",
                [
                    ["Milestone", "Target date", "Responsible"],
                    ["Roof lease signed with the municipality", "31 March 2026", "Hanna Vogt"],
                    ["Grid connection request sent", "20 March 2026", "Selin Aydın"],
                    ["Structural survey of both roofs", "25 April 2026", "Ingenieurbüro Kraft & Partner"],
                    ["Installer contract awarded", "30 April 2026", "Board"],
                    ["Installation (Whitsun holidays)", "26 May - 5 June 2026", "Lichtbau Solartechnik GmbH"],
                    ["Commissioning", "15 June 2026", "Lichtbau / Netze Mittelland"],
                ],
            ),
            ("h2", "Budget"),
            ("p", "Total budget 312,000 euros, financed by the KlimaKommunal grant (96,000 euros), a bank loan and new member shares."),
            ("h2", "Risks"),
            ("bullet", "Module delivery times (Helion has quoted 6 weeks)."),
            ("bullet", "Load capacity of the sports hall roof is unknown until the structural survey."),
            ("bullet", "The grid operator may require an upgrade of the local transformer station."),
        ],
    )

    # 5. Public flyer (HTML, DE, contains an error) -------------------------------
    write_text(
        target("public/2026-03-15-mitglieder-flyer.html"),
        """
        <!DOCTYPE html>
        <html lang="de">
        <head>
          <meta charset="utf-8">
          <title>Sonnendach Lindenhof - Jetzt Mitglied werden</title>
        </head>
        <body>
          <h1>Sonnendach Lindenhof: Sonnenstrom für unsere Schule</h1>
          <p>Die BürgerEnergie Eschenbrück eG baut eine Solaranlage mit 198 kWp auf
          der Grundschule am Lindenhof und der Sporthalle Nord.</p>
          <ul>
            <li>Inbetriebnahme: Juni 2026</li>
            <li>Batteriespeicher: 150 kWh</li>
            <li>Ein Anteil kostet nur 200 Euro</li>
            <li>Bis zu 40 Anteile pro Mitglied</li>
          </ul>
          <p>Infoabend: Donnerstag, 26. März 2026, 19 Uhr, Aula der Grundschule.</p>
          <p>Kontakt: info@eschenbrueck-energie.example</p>
          <footer>Stand: 15. März 2026</footer>
        </body>
        </html>
        """,
    )

    # 6. Grid connection request (EML) -------------------------------------------
    write_eml(
        target("emails/2026-03-18-netzanschlussanfrage.eml"),
        date="2026-03-18T09:12:00+01:00",
        sender=f"Selin Aydın <selin.aydin@{DOMAIN}>",
        to="Ute Sommer <u.sommer@netze-mittelland.example>",
        subject="Netzanschlussanfrage PV Grundschule am Lindenhof (198 kWp)",
        body="""
        Sehr geehrte Frau Sommer,

        die BürgerEnergie Eschenbrück eG beantragt den Netzanschluss für eine
        Photovoltaikanlage mit 198 kWp und einem Batteriespeicher mit 150 kWh
        an der Grundschule am Lindenhof, Lindenhofweg 2, Eschenbrück.

        Wir planen die Inbetriebnahme für den 15. Juni 2026. Die Unterlagen
        (Lageplan, Datenblätter Helion H-440, Wechselrichter Kestrel K-60)
        finden Sie im Anhang.

        Mit freundlichen Grüßen
        Selin Aydın
        Vorstand Technik, BürgerEnergie Eschenbrück eG
        Tel. 06543 555-0142
        """,
    )

    # 7. Structural survey (PDF, EN) ---------------------------------------------
    write_pdf(
        target("project/2026-04-22-structural-survey-summary.pdf"),
        "Structural survey summary - Sonnendach Lindenhof",
        [
            ("h1", "Structural survey: summary of findings"),
            ("p", "Client: BürgerEnergie Eschenbrück eG. Engineer: Dr. Ines Kraft, Ingenieurbüro Kraft & Partner. Date: 22 April 2026. Site visits: 14 and 16 April 2026."),
            ("h2", "Roof A - Grundschule am Lindenhof"),
            ("p", "Concrete flat roof, renovated in 2019. The load reserve of 0.85 kN/m2 is sufficient for a ballasted east-west or south-facing system. No restrictions."),
            ("h2", "Roof B - Sporthalle Nord"),
            ("p", "Timber roof structure from 1978. The west and middle sections have a load reserve of 0.42 kN/m2. The east section has a load reserve of only 0.12 kN/m2 and must not carry modules."),
            ("h2", "Recommendation"),
            ("p", "With the planned Helion H-440 modules (24.3 kg each) and ballast, the plant should be limited to approximately 165 kWp in total. Lighter modules would allow more modules on roof B; the client should send the data sheet of any alternative module for a re-check."),
            ("p", "The battery room in the school basement needs a fire door of class F90 before the battery is installed."),
            ("h2", "Fee"),
            ("p", "Invoice KP-2026-007 for this survey: 4,380 euros."),
        ],
    )

    # 8. Field notes (TXT, messy) ------------------------------------------------
    write_text(
        target("project/2026-04-30-selin-field-notes.txt"),
        """
        notes 30.4. -- call w/ Marco (Lichtbau) + Ines

        - statik: 165 kWp max w/ Helion (east part sporthalle = no go)
        - Ines: if modules lighter ~2 kg each -> could re-check, maybe ~400 modules ok??
        - Marco has contract ready, 220k for modules+inverters+battery+mounting
        - battery 150 kWh too expensive?? Tobias says budget tight, ask about 100 kWh
        - grid: no answer from Netze Mittelland yet (sent 18.3.!!) -> call Frau Sommer
        - F90 door for battery room: Jonas Feld (Bauamt) says municipality pays the door
        - scaffolding only in Pfingstferien, school says NO crane in the schoolyard during breaks

        TODO
        [ ] budget revision w/ Tobias before board mtg
        [ ] ask Marco about delivery times again, Helion "6 weeks" feels optimistic
        [x] send survey results to Hanna (89% yes!)
        """,
    )

    # 9. Site meeting transcript (TXT) -------------------------------------------
    write_text(
        target("meetings/2026-05-06-site-meeting-transcript.txt"),
        """
        Site meeting Sonnendach Lindenhof - transcript (auto-transcribed, lightly edited)
        Date: 6 May 2026, 15:30, school yard
        Speakers: Selin Aydın (BEE), Marco Petrović (Lichtbau Solartechnik), Jonas Feld (Bauamt), Dr. Ines Kraft (structural engineer)

        [00:00:12] Selin: Thanks for coming. The main points today are the crane position, the battery room and the delivery dates.
        [00:01:05] Marco: Crane can stand on the car park behind the sports hall. We need it for two days, Tuesday and Wednesday after Whitsun.
        [00:02:30] Jonas: The car park is fine. Please send the traffic plan to the Bauamt one week ahead.
        [00:03:48] Ines: For roof B, remember the east section stays empty. I marked it in red on the plan.
        [00:05:10] Selin: Marco, what is the status of the Helion modules?
        [00:05:22] Marco: Honestly, not good. Helion told us this morning that the H-440 now has a delivery time of 14 weeks. That would push us into September.
        [00:06:01] Selin: That is too late for the school holidays. Do you have an alternative?
        [00:06:15] Marco: Nordlicht has the NL-430 in stock. It is about 2.1 kg lighter per module. I will send an offer this week.
        [00:07:02] Ines: If the NL-430 is lighter, send me the data sheet. Roof B might take more modules then.
        [00:08:40] Jonas: The F90 fire door for the battery room will be installed by the municipality in the week of 18 May.
        [00:09:55] Selin: Good. Then the battery can come with the rest of the installation.
        [00:10:30] (end of recording)
        """,
    )

    # 10. Installer: module delay + substitution (EML thread) ----------------------
    write_eml(
        target("emails/2026-05-12-lichtbau-modulverzoegerung.eml"),
        date="2026-05-12T16:40:00+02:00",
        sender="Marco Petrović <m.petrovic@lichtbau-solar.example>",
        to=f"Selin Aydın <selin.aydin@{DOMAIN}>",
        subject="Module delivery Helion H-440 delayed - offer for Nordlicht NL-430",
        body="""
        Hi Selin,

        as discussed on site: Helion has confirmed a delivery time of 14 weeks
        for the H-440. We cannot install before the end of the Whitsun
        holidays with those modules.

        Alternative: Nordlicht NL-430 (430 Wp, 22.2 kg), available from stock
        in 3 weeks. Dr. Kraft has checked the data sheet: with the lighter
        module, roof A takes 276 modules and roof B takes 123 modules.
        That is 399 modules, i.e. 171.6 kWp in total.

        Price per Wp is almost the same; the module line in the contract
        becomes 399 x 118 EUR = 47,082 EUR.

        Timeline: installation 8-19 June, commissioning mid-July 2026 at the
        earliest (depends on Netze Mittelland).

        Please confirm by Wednesday so I can reserve the modules.

        Best regards
        Marco Petrović
        Project lead, Lichtbau Solartechnik GmbH
        """,
    )
    write_eml(
        target("emails/2026-05-13-re-modulverzoegerung.eml"),
        date="2026-05-13T08:05:00+02:00",
        sender=f"Selin Aydın <selin.aydin@{DOMAIN}>",
        to="Marco Petrović <m.petrovic@lichtbau-solar.example>",
        cc=f"Hanna Vogt <hanna.vogt@{DOMAIN}>, Tobias Brandt <tobias.brandt@{DOMAIN}>",
        subject="Re: Module delivery Helion H-440 delayed - offer for Nordlicht NL-430",
        in_reply_to="2026-05-12-lichtbau-modulverzoegerung",
        body="""
        Hi Marco,

        the board agreed last night: please go ahead with 399 x Nordlicht
        NL-430. The plant will therefore have 171.6 kWp instead of the
        198 kWp in the grant application; I will inform KlimaKommunal.

        We also reduce the battery from 150 kWh to 100 kWh (Tobias will send
        the revised budget). Commissioning target is now mid-July 2026.

        Thanks for finding a solution so quickly.

        Selin

        > Alternative: Nordlicht NL-430 (430 Wp, 22.2 kg), available from stock
        > in 3 weeks.
        """,
    )

    # 11. Budget revision (XLSX) --------------------------------------------------
    budget_rows: list[list[object]] = [
        ["Item", "Quantity", "Unit price (EUR)", "Total (EUR)", "Note"],
        ["PV modules Nordlicht NL-430", 399, 118, 47082, "replaces 450 x Helion H-440"],
        ["Inverters Kestrel K-60", 3, 6900, 20700, "WR-1, WR-2, WR-3"],
        ["Battery storage 100 kWh", 1, 54000, 54000, "reduced from 150 kWh"],
        ["Mounting and installation", 1, 98400, 98400, "Lichtbau Solartechnik GmbH"],
        ["Grid connection and transformer contribution", 1, 31500, 31500, "Netze Mittelland GmbH"],
        ["Structural survey", 1, 4380, 4380, "Ingenieurbüro Kraft & Partner"],
        ["Planning and permits", 1, 12600, 12600, ""],
        ["Monitoring system", 1, 3900, 3900, ""],
        ["Contingency", 1, 25938, 25938, "about 9.5 %"],
        ["Total", None, None, 298500, "was 312,000 in the grant application"],
    ]
    write_xlsx(
        target("finance/2026-05-18-budget-revision.xlsx"),
        {
            "Budget": budget_rows,
            "Financing": [
                ["Source", "Amount (EUR)", "Status"],
                ["KlimaKommunal grant", 96000, "approved 3 April 2026"],
                ["Loan Raiffeisenkasse Talgrund (3.1 %, 15 years)", 120000, "signed 11 May 2026"],
                ["New member shares (330 shares x 250 EUR)", 82500, "subscribed by 30 June 2026"],
                ["Total", 298500, ""],
            ],
            "Changes": [
                ["Change", "Before", "After", "Reason"],
                ["Plant size", "198 kWp", "171.6 kWp", "structural survey + module substitution"],
                ["Module", "Helion H-440", "Nordlicht NL-430", "14-week delivery delay"],
                ["Battery", "150 kWh", "100 kWh", "budget"],
                ["Total budget", "312,000 EUR", "298,500 EUR", ""],
                ["Expected yield", "190,000 kWh/year", "162,000 kWh/year", "smaller plant"],
            ],
        },
    )

    # 12. Member shares by year (CSV) --------------------------------------------
    write_csv(
        target("finance/2026-06-30-member-shares-by-year.csv"),
        [
            ["year", "new_members", "members_end_of_period", "shares_issued", "shares_end_of_period", "dividend_paid_pct"],
            ["2019", "61", "61", "1210", "1210", ""],
            ["2020", "84", "145", "1466", "2676", "0.0"],
            ["2021", "77", "222", "1204", "3880", "1.0"],
            ["2022", "69", "291", "902", "4782", "1.5"],
            ["2023", "41", "332", "520", "5302", "1.5"],
            ["2024", "31", "363", "380", "5682", "2.0"],
            ["2025", "25", "388", "230", "5912", "2.0"],
            ["2026-H1", "24", "412", "330", "6242", ""],
        ],
    )

    # 13. General assembly minutes (MD, DE) ---------------------------------------
    write_text(
        target("meetings/2026-06-13-generalversammlung-protokoll.md"),
        """
        # Protokoll der Generalversammlung 2026

        **BürgerEnergie Eschenbrück eG** · 13. Juni 2026, 18:00 Uhr · Aula der Grundschule am Lindenhof

        Versammlungsleitung: Klaus Reimann (Aufsichtsratsvorsitzender)
        Protokoll: Tobias Brandt

        Anwesend: 97 von 412 Mitgliedern, 23 weitere Mitglieder durch Vollmacht vertreten.

        ## TOP 1: Bericht des Vorstands

        Dr. Hanna Vogt berichtet über das Projekt Sonnendach Lindenhof. Wegen einer
        Lieferverzögerung von 14 Wochen wurden statt Helion H-440 nun 399 Module vom
        Typ Nordlicht NL-430 bestellt. Die Anlage hat damit 171,6 kWp. Der
        Batteriespeicher wurde aus Kostengründen auf 100 kWh verkleinert. Die
        Inbetriebnahme ist für Mitte Juli 2026 geplant.

        ## TOP 2: Jahresabschluss 2025

        Der Jahresabschluss 2025 wird einstimmig festgestellt. Vorstand und
        Aufsichtsrat werden entlastet.

        ## TOP 3: Dividende

        Beschluss: Für 2025 wird eine Dividende von 2,0 % ausgeschüttet
        (Auszahlung bis 30. Juni 2026). Ab dem Geschäftsjahr 2027 strebt die
        Genossenschaft eine Dividende von 3 % an.
        Abstimmung: 112 Ja, 5 Nein, 3 Enthaltungen.

        ## TOP 4: Projektbudget Sonnendach Lindenhof

        Beschluss: Die Generalversammlung genehmigt das überarbeitete Budget von
        298.500 Euro. Abstimmung: 116 Ja, 1 Nein, 3 Enthaltungen.

        ## TOP 5: Wahlen

        Selin Aydın wird für drei Jahre in den Vorstand wiedergewählt.

        ## TOP 6: Verschiedenes

        Die Wärmepumpe für das Freibad Eschenbrück soll im Herbst 2026 als
        Pilotprojekt geplant werden. Ein Mitglied fragt nach dem Anteilspreis auf
        dem Flyer; der Vorstand stellt klar, dass ein Anteil weiterhin 250 Euro
        kostet.
        """,
    )

    # 14. AGM slides (PPTX, EN) ---------------------------------------------------
    write_pptx(
        target("presentations/2026-06-13-agm-slides.pptx"),
        [
            ("General Assembly 2026", ["BürgerEnergie Eschenbrück eG - 13 June 2026"]),
            ("Our cooperative", ["412 members (388 at the end of 2025)", "6,242 shares at 250 euros each", "Dividend 2025: 2.0 %"]),
            (
                "Sonnendach Lindenhof: status",
                [
                    "171.6 kWp: 399 Nordlicht NL-430 modules (276 school roof, 123 sports hall)",
                    "Battery storage: 100 kWh in the school basement",
                    "Installation: 8-19 June 2026",
                    "Commissioning: mid-July 2026",
                    "Expected yield: 162,000 kWh per year, about 62 tonnes of CO2 saved",
                ],
            ),
            (
                "Financing: 298,500 euros",
                ["KlimaKommunal grant: 96,000 euros", "Raiffeisenkasse Talgrund loan: 120,000 euros", "New member shares: 82,500 euros"],
            ),
            ("Dividend policy", ["2025: 2.0 % (proposal)", "Target from 2027: 3 %"]),
            ("Next project: Freibad heat pump", ["60 kW heat pump for the outdoor pool", "Pilot planned for autumn 2026", "71 % of surveyed members in favour"]),
        ],
    )

    # 15. Grid approval (EML, DE) -------------------------------------------------
    write_eml(
        target("emails/2026-06-24-netzanschlusszusage.eml"),
        date="2026-06-24T11:30:00+02:00",
        sender="Ute Sommer <u.sommer@netze-mittelland.example>",
        to=f"Selin Aydın <selin.aydin@{DOMAIN}>",
        subject="Netzanschlusszusage PV Grundschule am Lindenhof",
        body="""
        Sehr geehrte Frau Aydın,

        wir bestätigen den Netzanschluss für Ihre Photovoltaikanlage mit
        171,6 kWp (geänderte Anlagengröße vom 14. Mai 2026).

        Die Ortsnetzstation Lindenhof muss dafür ertüchtigt werden. Die
        Zuschaltung der Anlage ist deshalb frühestens ab dem 1. August 2026
        möglich. Den Baukostenzuschuss von 31.500 Euro stellen wir Ihnen
        gesondert in Rechnung.

        Mit freundlichen Grüßen
        Ute Sommer
        Netzanschlüsse, Netze Mittelland GmbH
        """,
    )

    # 16. Board minutes, July (MD, EN) --------------------------------------------
    write_text(
        target("meetings/2026-07-15-board-minutes.md"),
        """
        # Board meeting minutes - 15 July 2026

        BürgerEnergie Eschenbrück eG. Present: Hanna Vogt, Tobias Brandt, Selin Aydın.

        ## Commissioning date

        Installation finished on 19 June 2026. Netze Mittelland can only connect
        the plant from 1 August 2026 because the local transformer station must be
        upgraded first. New commissioning target: 1 August 2026.

        ## Membership

        Membership is now 419 (412 at the general assembly, plus seven new members
        who joined through the school's parent association).

        ## Freibad heat pump pilot

        The outdoor pool's electrical connection would need an upgrade costing about
        18,000 euros, which is not in any budget. Decision: the heat pump pilot is
        postponed to spring 2027. Hanna will inform the members.

        ## Flyer

        The March flyer said a share costs 200 euros; the correct price is 250 euros.
        Tobias has corrected the reprint and answered the members who asked.
        """,
    )

    # 17. Member question + reply (EML) -------------------------------------------
    write_eml(
        target("emails/2026-07-08-mitgliederfrage-anteil.eml"),
        date="2026-07-08T20:14:00+02:00",
        sender="Rainer Holm <rainer.holm@mailbox.example>",
        to=f"info@{DOMAIN}",
        subject="Frage zum Anteilspreis und zur Dividende",
        body="""
        Hallo zusammen,

        auf dem Flyer vom März steht, ein Anteil kostet 200 Euro, auf eurer
        Website stehen 250 Euro. Was stimmt denn? Ich möchte für meine
        Tochter 4 Anteile zeichnen.

        Und wann wird die Dividende für 2025 ausgezahlt?

        Viele Grüße
        Rainer Holm
        """,
    )
    write_eml(
        target("emails/2026-07-09-re-mitgliederfrage-anteil.eml"),
        date="2026-07-09T09:02:00+02:00",
        sender=f"Tobias Brandt <tobias.brandt@{DOMAIN}>",
        to="Rainer Holm <rainer.holm@mailbox.example>",
        subject="Re: Frage zum Anteilspreis und zur Dividende",
        in_reply_to="2026-07-08-mitgliederfrage-anteil",
        body="""
        Hallo Herr Holm,

        danke für den Hinweis: Ein Anteil kostet 250 Euro. Die 200 Euro auf dem
        Flyer vom März sind ein Druckfehler, der Nachdruck ist korrigiert.
        Vier Anteile kosten also 1.000 Euro.

        Die Dividende von 2,0 % für 2025 wurde am 30. Juni 2026 ausgezahlt.
        Ab 2027 streben wir 3 % an.

        Viele Grüße
        Tobias Brandt
        Vorstand Finanzen
        """,
    )

    # 18. FAQ (MD, EN, one stale fact) --------------------------------------------
    write_text(
        target("public/faq.md"),
        """
        # Sonnendach Lindenhof - frequently asked questions

        *Last updated: 20 July 2026*

        **Who builds the plant?**
        BürgerEnergie Eschenbrück eG, our citizens' energy cooperative, together with
        the installer Lichtbau Solartechnik GmbH.

        **How big is the plant?**
        171.6 kWp on the roofs of the Grundschule am Lindenhof and the Sporthalle Nord.

        **Is there a battery?**
        Yes, a 150 kWh battery in the school basement stores midday power for the afternoon.

        **When will it produce electricity?**
        From August 2026, as soon as Netze Mittelland has upgraded the local transformer station.

        **What does a share cost?**
        250 euros. Each member can hold up to 40 shares.

        **Who uses the electricity?**
        The school buys it at 19.5 cents per kWh for 20 years; the surplus goes into the grid.

        **What about the pool heat pump?**
        The pilot has been postponed to spring 2027.
        """,
    )

    # 19. Commissioning report (EML) ----------------------------------------------
    write_eml(
        target("emails/2026-08-19-inbetriebnahme-bericht.eml"),
        date="2026-08-19T17:55:00+02:00",
        sender="Marco Petrović <m.petrovic@lichtbau-solar.example>",
        to=f"Selin Aydın <selin.aydin@{DOMAIN}>",
        cc=f"Hanna Vogt <hanna.vogt@{DOMAIN}>",
        subject="Commissioning report Sonnendach Lindenhof - plant is live",
        body="""
        Hi Selin,

        the plant was energised today, 19 August 2026, at 10:42 after
        Netze Mittelland switched the new transformer station on.

        Summary:
        - 399 x Nordlicht NL-430, 171.6 kWp
        - 3 inverters Kestrel K-60 (WR-1, WR-2, WR-3), all online
        - battery storage 100 kWh, first charge completed at 14:10
        - before energising we replaced a faulty string fuse on string 7 (WR-3)

        The first-day yield until 17:30 was 612 kWh. The acceptance protocol
        follows by post.

        Best regards
        Marco
        """,
    )

    # 20. Newspaper article (MD, DE) ----------------------------------------------
    write_text(
        target("press/2026-08-22-eschenbruecker-anzeiger.md"),
        """
        # Sonnenstrom für den Lindenhof

        *Eschenbrücker Anzeiger, 22. August 2026*

        Die Grundschule am Lindenhof bekommt ihren Strom jetzt vom eigenen Dach.
        Am 19. August hat die BürgerEnergie Eschenbrück eG ihre neue Solaranlage in
        Betrieb genommen. Die 399 Module auf der Schule und der Sporthalle Nord
        leisten zusammen 171,6 Kilowatt Peak.

        "Das ist ein Projekt von Eschenbrückern für Eschenbrück", sagte die
        Vorstandsvorsitzende Dr. Hanna Vogt. Die Genossenschaft hat inzwischen über
        400 Mitglieder. Bürgermeisterin Petra Lindqvist lobte die Zusammenarbeit:
        Die Gemeinde verpachtet die Dächer für einen symbolischen Euro im Jahr.

        Ein Batteriespeicher im Keller sorgt dafür, dass die Schule auch
        nachmittags Solarstrom nutzen kann. Die geplante Wärmepumpe für das
        Freibad soll im Frühjahr 2027 folgen.
        """,
    )

    # 21. Monitoring: daily production (CSV) --------------------------------------
    daily = [
        ("2026-08-19", 612, 0, 540, 72),
        ("2026-08-20", 948, 118, 702, 128),
        ("2026-08-21", 1004, 121, 745, 138),
        ("2026-08-22", 655, 96, 470, 89),
        ("2026-08-23", 402, 60, 290, 52),
        ("2026-08-24", 877, 115, 640, 122),
        ("2026-08-25", 991, 124, 731, 136),
        ("2026-08-26", 1032, 126, 764, 142),
        ("2026-08-27", 1061, 129, 788, 144),
        ("2026-08-28", 812, 110, 590, 112),
        ("2026-08-29", 544, 80, 395, 69),
        ("2026-08-30", 931, 112, 690, 129),
        ("2026-08-31", 968, 119, 715, 134),
        ("2026-09-01", 902, 117, 661, 124),
        ("2026-09-02", 715, 102, 518, 95),
        ("2026-09-03", 488, 74, 351, 63),
        ("2026-09-04", 866, 113, 634, 119),
        ("2026-09-05", 897, 115, 657, 125),
        ("2026-09-06", 823, 108, 601, 114),
        ("2026-09-07", 855, 402, 283, 170),
        ("2026-09-08", 839, 411, 262, 166),
        ("2026-09-09", 596, 355, 118, 123),
        ("2026-09-10", 781, 398, 222, 161),
        ("2026-09-11", 802, 405, 234, 163),
        ("2026-09-12", 744, 98, 541, 105),
        ("2026-09-13", 698, 93, 507, 98),
        ("2026-09-14", 772, 399, 212, 161),
    ]
    rows: list[list[object]] = [["date", "pv_kwh", "school_consumption_kwh", "grid_feed_in_kwh", "battery_charged_kwh"]]
    rows += [list(d) for d in daily]
    write_csv(target("monitoring/2026-09-pv-daily-production.csv"), rows)

    # 22. Monitoring: inverter status (JSON) --------------------------------------
    status = {
        "plant": "Sonnendach Lindenhof",
        "operator": "BürgerEnergie Eschenbrück eG",
        "snapshot_utc": "2026-09-15T06:00:00Z",
        "peak_power_kwp": 171.6,
        "energy_since_commissioning_kwh": sum(d[1] for d in daily),
        "inverters": [
            {"id": "WR-1", "model": "Kestrel K-60", "state": "ok", "strings": 8, "energy_kwh": 7641},
            {"id": "WR-2", "model": "Kestrel K-60", "state": "ok", "strings": 8, "energy_kwh": 7598},
            {
                "id": "WR-3",
                "model": "Kestrel K-60",
                "state": "warning",
                "strings": 7,
                "energy_kwh": sum(d[1] for d in daily) - 7641 - 7598,
                "active_warnings": [
                    {
                        "code": "ISO-217",
                        "text": "Isolation resistance low on string 7 (180 kOhm, limit 500 kOhm)",
                        "first_seen": "2026-09-03T05:58:00Z",
                        "occurrences": 6,
                    }
                ],
            },
        ],
        "battery": {"capacity_kwh": 100, "state_of_charge_pct": 18, "cycles": 24},
    }
    target("monitoring/2026-09-15-inverter-status.json").write_text(
        json.dumps(status, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    # 23. Monitoring: data logger (LOG) -------------------------------------------
    write_text(
        target("monitoring/2026-09-14-datalogger.log"),
        """
        2026-09-14T04:30:02Z INFO  logger started, firmware 2.8.1, 3 inverters configured
        2026-09-14T05:41:17Z INFO  WR-1 feed-in started
        2026-09-14T05:41:40Z INFO  WR-2 feed-in started
        2026-09-14T05:44:03Z WARN  WR-3 ISO-217 isolation resistance low on string 7 (180 kOhm)
        2026-09-14T06:12:55Z INFO  WR-3 ISO-217 cleared, feed-in started
        2026-09-14T09:00:00Z INFO  battery SoC 41 %, charging 22.4 kW
        2026-09-14T12:00:00Z INFO  plant output 121.8 kW, school load 38.1 kW
        2026-09-14T15:00:00Z INFO  battery SoC 100 %, charging stopped
        2026-09-14T18:59:31Z INFO  all inverters in night mode
        2026-09-14T19:00:00Z INFO  daily yield 772 kWh
        """,
    )

    # 24. Monitoring config (YAML) ------------------------------------------------
    write_text(
        target("monitoring/site-config.yaml"),
        """
        # Monitoring configuration - Sonnendach Lindenhof
        plant:
          name: Sonnendach Lindenhof
          operator: BürgerEnergie Eschenbrück eG
          peak_power_kwp: 171.6
          commissioned: 2026-08-19
        modules:
          type: Nordlicht NL-430
          count: 399
          roofs:
            school_roof_a: 276
            sports_hall_roof_b: 123
        inverters:
          - id: WR-1
            model: Kestrel K-60
            roof: school_roof_a
          - id: WR-2
            model: Kestrel K-60
            roof: school_roof_a
          - id: WR-3
            model: Kestrel K-60
            roof: sports_hall_roof_b
        battery:
          capacity_kwh: 100
          location: school basement
        alerts:
          email: technik@eschenbrueck-energie.example
          isolation_warning_threshold_kohm: 500
        """,
    )

    # 25. Battery events (TSV) ----------------------------------------------------
    write_csv(
        target("monitoring/2026-09-battery-events.tsv"),
        [
            ["timestamp", "event", "soc_pct", "note"],
            ["2026-08-19T14:10:00", "first_full_charge", "100", "commissioning day"],
            ["2026-09-01T02:00:00", "firmware_update", "35", "battery management 1.4.2"],
            ["2026-09-07T07:45:00", "discharge_start_school_day", "96", "first school day after holidays"],
            ["2026-09-09T11:20:00", "charge_limited", "62", "cloudy, low PV output"],
            ["2026-09-14T15:00:00", "full_charge", "100", ""],
        ],
        delimiter="\t",
    )

    # 26. Roof layout (PNG) -------------------------------------------------------
    write_png(target("media/2026-05-lindenhof-roof-layout.png"))

    # 27. Invoice archive (ZIP) ---------------------------------------------------
    invoices = {
        "invoice-LB-2026-031-deposit.txt": textwrap.dedent(
            """
            Lichtbau Solartechnik GmbH - Invoice LB-2026-031
            Date: 20 May 2026
            Customer: BürgerEnergie Eschenbrück eG
            Project: Sonnendach Lindenhof
            Deposit 20 % of contract value 220,182.00 EUR: 44,036.40 EUR
            Payment due: 3 June 2026
            """
        ).lstrip(),
        "invoice-KP-2026-007-structural-survey.txt": textwrap.dedent(
            """
            Ingenieurbüro Kraft & Partner - Invoice KP-2026-007
            Date: 24 April 2026
            Structural survey Grundschule am Lindenhof and Sporthalle Nord
            Amount: 4,380.00 EUR
            """
        ).lstrip(),
        "invoice-NM-2026-4471-grid.txt": textwrap.dedent(
            """
            Netze Mittelland GmbH - Invoice NM-2026-4471
            Date: 1 July 2026
            Construction cost contribution, transformer station Lindenhof: 31,500.00 EUR
            """
        ).lstrip(),
        "payments-2026.csv": "date,payee,invoice,amount_eur\n"
        "2026-04-30,Ingenieurbüro Kraft & Partner,KP-2026-007,4380.00\n"
        "2026-06-02,Lichtbau Solartechnik GmbH,LB-2026-031,44036.40\n"
        "2026-07-14,Netze Mittelland GmbH,NM-2026-4471,31500.00\n",
    }
    archive = target("archive/2026-08-invoices.zip")
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in invoices.items():
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, text.encode("utf-8"))

    # 28. Roof lease summary (TXT, DE) --------------------------------------------
    write_text(
        target("project/2026-03-26-dachnutzungsvertrag-zusammenfassung.txt"),
        """
        Dachnutzungsvertrag - Zusammenfassung (intern)

        Vertragspartner: Gemeinde Eschenbrück, vertreten durch Bürgermeisterin
        Petra Lindqvist, und BürgerEnergie Eschenbrück eG, vertreten durch
        Dr. Hanna Vogt und Tobias Brandt.
        Unterzeichnet am 26. März 2026.

        - Dächer: Grundschule am Lindenhof und Sporthalle Nord
        - Laufzeit: 20 Jahre, Verlängerungsoption um 5 Jahre
        - Pacht: 1 Euro pro Jahr (symbolisch)
        - Stromlieferung: Die Gemeinde kauft den Solarstrom für die Schule zu
          19,5 ct/kWh (Stromliefervertrag, ebenfalls 20 Jahre)
        - Die Gemeinde übernimmt den Einbau der F90-Brandschutztür im Batterieraum
        - Rückbau nach Vertragsende auf Kosten der Genossenschaft
        """,
    )

    return written


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=RAW_DIR, help="output folder (default: data/raw)")
    parser.add_argument("--clean", action="store_true", help="delete everything in the output folder first")
    args = parser.parse_args(argv)
    out: Path = args.out
    if args.clean and out.exists():
        for child in out.iterdir():
            if child.name.startswith("."):
                continue
            shutil.rmtree(child) if child.is_dir() else child.unlink()
    out.mkdir(parents=True, exist_ok=True)
    written = build(out)
    for path in written:
        print(path.relative_to(out))
    print(f"\n{len(written)} files written to {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
